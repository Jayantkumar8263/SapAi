"""
BSP Inventory AI Agent - Tool Layer

This module contains the tools that the AI agent can use
to execute the BSP inventory workflow.

IMPORTANT:
- This file does NOT directly control SAP yet.
- SAP integration will be added later.
- Existing deterministic Python services are reused here.
- Passwords are never printed or returned.
"""
from app.services.bsp_workbook_builder import (
    build_bsp_workbook,
)

from pathlib import Path
from typing import Any, Dict, Optional

import pandas as pd

from app.agent.state import AgentState
from app.services.inventory_processor import process_inventory_file
from app.services.inventory_comparator import compare_inventory
from app.services.inventory_ai import analyze_inventory_changes
from app.services.inventory_classifier import classify_inventory


# ============================================================
# PROJECT PATHS
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parents[3]

DATA_DIR = PROJECT_ROOT / "data"

REPORT_DIR = DATA_DIR / "reports"

REPORT_DIR.mkdir(
    parents=True,
    exist_ok=True,
)


# ============================================================
# TOOL RESULT HELPER
# ============================================================

def tool_result(
    success: bool,
    message: str,
    **kwargs: Any,
) -> Dict[str, Any]:
    """
    Standard response format for agent tools.
    """

    result = {
        "success": success,
        "message": message,
    }

    result.update(kwargs)

    return result


# ============================================================
# 1. GET AGENT CONFIGURATION
# ============================================================

def get_agent_configuration(
    state: AgentState,
) -> Dict[str, Any]:
    """
    Return the configuration currently being used by the agent.

    Password is deliberately excluded.
    """

    return tool_result(
        True,
        "Agent configuration loaded.",
        configuration={
            "username": state.username,
            "tcode": state.tcode,
            "variant": state.variant,
            "plant": state.plant,
            "previous_period": state.previous_period,
            "current_period": state.current_period,
            "preferences": state.preferences,
        },
    )


# ============================================================
# 2. VALIDATE SAP CREDENTIAL REQUIREMENTS
# ============================================================

def validate_sap_credentials(
    state: AgentState,
) -> Dict[str, Any]:
    """
    Check whether the information required to connect to SAP
    has been supplied.

    This does NOT connect to SAP.
    """

    missing = []

    if not state.username:
        missing.append("SAP username")

    if not state.password:
        missing.append("SAP password")

    if missing:
        return tool_result(
            False,
            "SAP credentials are incomplete.",
            missing=missing,
        )

    return tool_result(
        True,
        "SAP credentials are available for the session.",
        username=state.username,
    )


# ============================================================
# 3. SAP LOGIN - PLACEHOLDER
# ============================================================

def login_to_sap(
    state: AgentState,
    simulation: bool = True,
) -> Dict[str, Any]:
    """
    SAP login tool.

    Current version:
        Simulation only.

    Future version:
        Real SAP integration.
    """

    credential_check = validate_sap_credentials(state)

    if not credential_check["success"]:
        return credential_check

    if simulation:
        state.status = "running"
        state.mark_completed("sap_login")

        return tool_result(
            True,
            "SAP login simulated successfully.",
            mode="simulation",
            username=state.username,
        )

    return tool_result(
        False,
        "Real SAP connection is not implemented yet.",
        mode="real",
    )


# ============================================================
# 4. PREPARE SAP WORKFLOW PARAMETERS
# ============================================================

def prepare_sap_workflow(
    state: AgentState,
) -> Dict[str, Any]:
    """
    Prepare parameters required for the BSP inventory workflow.

    Default:
        T-Code  -> MC.1
        Variant -> B002159
    """

    if not state.tcode:
        state.tcode = "MC.1"

    if not state.variant:
        state.variant = "B002159"

    state.mark_completed(
        "workflow_parameters_prepared"
    )

    return tool_result(
        True,
        "SAP workflow parameters prepared.",
        parameters={
            "tcode": state.tcode,
            "variant": state.variant,
            "plant": state.plant,
            "previous_period": state.previous_period,
            "current_period": state.current_period,
        },
    )


# ============================================================
# 5. RUN MC.1 - PLACEHOLDER
# ============================================================

def run_mc1(
    state: AgentState,
    simulation: bool = True,
) -> Dict[str, Any]:
    """
    Run SAP transaction MC.1.
    """

    if state.tcode != "MC.1":
        return tool_result(
            False,
            f"Unexpected T-code: {state.tcode}. Expected MC.1.",
        )

    if simulation:
        state.mark_completed("mc1_executed")

        return tool_result(
            True,
            "MC.1 execution simulated successfully.",
            tcode="MC.1",
            mode="simulation",
        )

    return tool_result(
        False,
        "Real MC.1 execution is not implemented yet.",
        mode="real",
    )


# ============================================================
# 6. SELECT BSP VARIANT
# ============================================================

def select_variant(
    state: AgentState,
    simulation: bool = True,
) -> Dict[str, Any]:
    """
    Select the BSP inventory variant.
    """

    if not state.variant:
        return tool_result(
            False,
            "No SAP variant has been configured.",
        )

    if simulation:
        state.mark_completed("variant_selected")

        return tool_result(
            True,
            "SAP variant selected successfully.",
            variant=state.variant,
            mode="simulation",
        )

    return tool_result(
        False,
        "Real SAP variant selection is not implemented yet.",
        mode="real",
    )


# ============================================================
# 7. SELECT REPORT PERIOD
# ============================================================

def select_report_period(
    state: AgentState,
    previous_period: Optional[str] = None,
    current_period: Optional[str] = None,
    simulation: bool = True,
) -> Dict[str, Any]:
    """
    Configure previous and current reporting periods.
    """

    if previous_period:
        state.previous_period = previous_period

    if current_period:
        state.current_period = current_period

    if not state.previous_period:
        return tool_result(
            False,
            "Previous reporting period is required.",
        )

    if not state.current_period:
        return tool_result(
            False,
            "Current reporting period is required.",
        )

    if simulation:
        state.mark_completed("report_period_selected")

        return tool_result(
            True,
            "Reporting periods configured successfully.",
            previous_period=state.previous_period,
            current_period=state.current_period,
            mode="simulation",
        )

    return tool_result(
        False,
        "Real SAP period selection is not implemented yet.",
        mode="real",
    )


# ============================================================
# 8. RUN PLANT ANALYSIS
# ============================================================

def run_plant_analysis(
    state: AgentState,
    simulation: bool = True,
) -> Dict[str, Any]:
    """
    Run Plant Analysis inside SAP.
    """

    if simulation:
        state.mark_completed("plant_analysis")

        return tool_result(
            True,
            "Plant Analysis simulated successfully.",
            plant=state.plant,
            mode="simulation",
        )

    return tool_result(
        False,
        "Real SAP Plant Analysis is not implemented yet.",
        mode="real",
    )


# ============================================================
# 9. EXPORT INVENTORY FROM SAP - PLACEHOLDER
# ============================================================

def export_inventory_from_sap(
    state: AgentState,
    simulation: bool = True,
) -> Dict[str, Any]:
    """
    Export inventory data.

    Simulation mode uses local inventory files.
    """

    if simulation:

        previous_file = DATA_DIR / "previous_inventory.xlsx"

        current_file = DATA_DIR / "current_inventory.xlsx"

        if not previous_file.exists():
            return tool_result(
                False,
                "Previous inventory file not found.",
                expected_file=str(previous_file),
            )

        if not current_file.exists():
            return tool_result(
                False,
                "Current inventory file not found.",
                expected_file=str(current_file),
            )

        state.previous_inventory_file = str(
            previous_file
        )

        state.current_inventory_file = str(
            current_file
        )

        state.mark_completed(
            "sap_inventory_export"
        )

        return tool_result(
            True,
            "Local inventory files loaded as simulated SAP exports.",
            previous_file=str(previous_file),
            current_file=str(current_file),
            mode="simulation",
        )

    return tool_result(
        False,
        "Real SAP inventory export is not implemented yet.",
        mode="real",
    )


# ============================================================
# 10. PROCESS INVENTORY FILE
# ============================================================

def process_inventory(
    file_path: str,
) -> Dict[str, Any]:
    """
    Process a single SAP inventory Excel file.
    """

    path = Path(file_path)

    if not path.exists():
        return tool_result(
            False,
            "Inventory file does not exist.",
            file_path=str(path),
        )

    try:

        result = process_inventory_file(
            str(path)
        )

        if not result.get("success", False):

            return tool_result(
                False,
                "Inventory processing failed validation.",
                file_path=str(path),
                validation=result.get(
                    "validation"
                ),
            )

        return tool_result(
            True,
            "Inventory file processed successfully.",
            file_path=str(path),
            validation=result.get(
                "validation",
                {},
            ),
            summary=result.get(
                "summary",
                {},
            ),
            data=result.get(
                "data"
            ),
        )

    except Exception as exc:

        return tool_result(
            False,
            f"Inventory processing failed: {exc}",
            file_path=str(path),
        )


# ============================================================
# 11. PROCESS PREVIOUS INVENTORY
# ============================================================

def process_previous_inventory(
    state: AgentState,
) -> Dict[str, Any]:
    """
    Process previous-month inventory.
    """

    if not state.previous_inventory_file:

        return tool_result(
            False,
            "Previous inventory file has not been provided.",
        )

    result = process_inventory(
        state.previous_inventory_file
    )

    if result["success"]:

        state.previous_inventory = result.get(
            "data"
        )

        state.mark_completed(
            "previous_inventory_processed"
        )

    return result


# ============================================================
# 12. PROCESS CURRENT INVENTORY
# ============================================================

def process_current_inventory(
    state: AgentState,
) -> Dict[str, Any]:
    """
    Process current-month inventory.
    """

    if not state.current_inventory_file:

        return tool_result(
            False,
            "Current inventory file has not been provided.",
        )

    result = process_inventory(
        state.current_inventory_file
    )

    if result["success"]:

        state.current_inventory = result.get(
            "data"
        )

        state.mark_completed(
            "current_inventory_processed"
        )

    return result


# ============================================================
# 13. COMPARE INVENTORY
# ============================================================

def compare_inventory_files(
    state: AgentState,
) -> Dict[str, Any]:
    """
    Compare previous and current inventory.
    """

    if not state.previous_inventory_file:

        return tool_result(
            False,
            "Previous inventory file is missing.",
        )

    if not state.current_inventory_file:

        return tool_result(
            False,
            "Current inventory file is missing.",
        )

    previous_path = Path(
        state.previous_inventory_file
    )

    current_path = Path(
        state.current_inventory_file
    )

    if not previous_path.exists():

        return tool_result(
            False,
            "Previous inventory file does not exist.",
            file_path=str(previous_path),
        )

    if not current_path.exists():

        return tool_result(
            False,
            "Current inventory file does not exist.",
            file_path=str(current_path),
        )

    try:

        comparison, summary = compare_inventory(
            str(previous_path),
            str(current_path),
        )

        state.comparison = comparison

        state.summary = clean_summary(
            summary
        )

        state.mark_completed(
            "inventory_comparison"
        )

        return tool_result(
            True,
            "Inventory comparison completed successfully.",
            summary=state.summary,
            comparison=comparison,
        )

    except Exception as exc:

        state.add_error(
            f"Inventory comparison failed: {exc}"
        )

        return tool_result(
            False,
            f"Inventory comparison failed: {exc}",
        )


# ============================================================
# 14. AI INVENTORY ANALYSIS
# ============================================================

def analyze_inventory(
    state: AgentState,
) -> Dict[str, Any]:
    """
    Analyze inventory changes after deterministic comparison.

    Deterministic comparison remains the source of truth.
    """

    comparison = getattr(
        state,
        "comparison",
        None,
    )

    if comparison is None:

        comparison = getattr(
            state,
            "comparison_result",
            None,
        )

    if comparison is None:

        return tool_result(
            False,
            "Inventory comparison must be completed before AI analysis.",
        )

    summary = getattr(
        state,
        "summary",
        None,
    )

    if summary is None:
        summary = {}

    try:

        result = analyze_inventory_changes(
            comparison=comparison,
            summary=summary,
        )

        if not result.get("success"):

            return tool_result(
                False,
                result.get(
                    "message",
                    "Inventory analysis failed.",
                ),
            )

        analysis = result.get(
            "analysis",
            {},
        )

        if hasattr(
            state,
            "results",
        ):

            state.results[
                "inventory_analysis"
            ] = analysis

        try:

            state.inventory_analysis = analysis

        except Exception:

            pass

        state.mark_completed(
            "inventory_analysis"
        )

        return tool_result(
            True,
            "Inventory intelligence analysis completed successfully.",
            analysis=analysis,
        )

    except Exception as exc:

        try:

            state.add_error(
                f"Inventory analysis failed: {exc}"
            )

        except Exception:

            pass

        return tool_result(
            False,
            f"Inventory analysis failed: {exc}",
        )


# ============================================================
# 15. CLEAN INVENTORY SUMMARY
# ============================================================

def clean_summary(
    summary: Any,
) -> Dict[str, Any]:
    """
    Convert pandas / NumPy values into JSON-safe Python values.
    """

    if summary is None:
        return {}

    if not isinstance(summary, dict):
        return {
            "value": summary
        }

    cleaned = {}

    for key, value in summary.items():

        if hasattr(
            value,
            "item",
        ):

            try:

                value = value.item()

            except Exception:

                pass

        if isinstance(
            value,
            float,
        ):

            if pd.isna(value):
                value = None

        cleaned[key] = value

    return cleaned


# ============================================================
# 16. GET COMPARISON SUMMARY
# ============================================================

def get_comparison_summary(
    state: AgentState,
) -> Dict[str, Any]:
    """
    Return the latest inventory comparison summary.
    """

    if not state.summary:

        return tool_result(
            False,
            "No inventory comparison has been completed.",
        )

    return tool_result(
        True,
        "Comparison summary retrieved.",
        summary=state.summary,
    )

def get_bsp_reference_file(state: AgentState) -> Path:
    """
    Return the RI/Capital reference workbook used by
    the BSP inventory process.
    """

    configured = (
        state.preferences.get("bsp_reference_file")
        if isinstance(state.preferences, dict)
        else None
    )

    if not configured:
        raise FileNotFoundError(
            "BSP RI/Capital reference workbook is not configured. "
            "Set state.preferences['bsp_reference_file']."
        )

    reference_file = Path(configured)

    if not reference_file.exists():
        raise FileNotFoundError(
            f"BSP RI/Capital reference workbook was not found: "
            f"{reference_file}"
        )

    return reference_file

def generate_bsp_inventory_workbook(
    state: AgentState,
) -> Dict[str, Any]:
    """
    Generate the actual BSP five-sheet inventory workbook.

    Sheets:
        Inventory Prev
        Inventory Pres
        RI
        Capital
        Inventory
    """

    if not state.previous_inventory_file:
        return tool_result(
            False,
            "Previous inventory file is required.",
        )

    if not state.current_inventory_file:
        return tool_result(
            False,
            "Current inventory file is required.",
        )

    try:
        reference_file = get_bsp_reference_file(state)

        output_file = (
            REPORT_DIR
            / "BSP_Final_Inventory.xlsx"
        )

        result = build_bsp_workbook(
            previous_file=state.previous_inventory_file,
            present_file=state.current_inventory_file,
            reference_file=reference_file,
            output_file=output_file,
        )

        if not result.get("success", False):
            return tool_result(
                False,
                result.get(
                    "message",
                    "BSP workbook generation failed.",
                ),
            )

        state.report_file = str(
            result["output_file"]
        )

        state.report_path = str(
            result["output_file"]
        )

        state.add_result(
            "bsp_inventory_workbook",
            result,
        )

        state.metadata[
            "bsp_inventory_workbook"
        ] = result

        state.mark_completed(
            "bsp_inventory_workbook_generated"
        )

        return tool_result(
            True,
            "BSP final inventory workbook generated successfully.",
            workbook=result["output_file"],
            final_inventory_rows=result.get(
                "final_inventory_rows",
                0,
            ),
            previous_rows=result.get(
                "previous_rows",
                0,
            ),
            present_rows=result.get(
                "present_rows",
                0,
            ),
            ri_materials=result.get(
                "ri_materials",
                0,
            ),
            capital_materials=result.get(
                "capital_materials",
                0,
            ),
        )

    except Exception as exc:
        state.add_error(
            f"BSP workbook generation failed: {exc}"
        )

        return tool_result(
            False,
            f"BSP workbook generation failed: {exc}",
        )
# ============================================================
# 17. GENERATE REPORT
# ============================================================

def generate_report(
    state: AgentState,
) -> Dict[str, Any]:
    """
    Generate the final inventory report.
    """

    if state.comparison is None:

        return tool_result(
            False,
            "Inventory comparison must be completed before generating the report.",
        )

    try:

        from app.services.inventory_report import (
            generate_inventory_report,
        )

        output_file = (
            REPORT_DIR
            / "BSP_Inventory_Report.xlsx"
        )

        report = generate_inventory_report(
            state.comparison,
            state.summary,
            str(output_file),
        )

        if report:

            state.report_file = str(
                report
            )

        else:

            state.report_file = str(
                output_file
            )

        state.mark_completed(
            "report_generated"
        )

        return tool_result(
            True,
            "Inventory report generated successfully.",
            report=state.report_file,
        )

    except TypeError:

        try:

            report = generate_inventory_report()

            state.report_file = str(
                report
            )

            state.mark_completed(
                "report_generated"
            )

            return tool_result(
                True,
                "Inventory report generated successfully.",
                report=state.report_file,
            )

        except Exception as exc:

            state.add_error(
                f"Report generation failed: {exc}"
            )

            return tool_result(
                False,
                f"Report generation failed: {exc}",
            )

    except Exception as exc:

        state.add_error(
            f"Report generation failed: {exc}"
        )

        return tool_result(
            False,
            f"Report generation failed: {exc}",
        )


# ============================================================
# 18. GET REPORT PATH
# ============================================================

def get_report_path(
    state: AgentState,
) -> Dict[str, Any]:
    """
    Return generated report location.
    """

    if not state.report_file:

        return tool_result(
            False,
            "No report has been generated yet.",
        )

    report_path = Path(
        state.report_file
    )

    return tool_result(
        True,
        "Report is available.",
        report=str(report_path),
        exists=report_path.exists(),
    )


# ============================================================
# 19. WORKFLOW STATUS
# ============================================================

def get_workflow_status(
    state: AgentState,
) -> Dict[str, Any]:
    """
    Return current workflow state.

    Password is never returned.
    """

    return tool_result(
        True,
        "Workflow status retrieved.",
        state=state.safe_dict(),
    )


# ============================================================
# 20. RESET WORKFLOW
# ============================================================

def reset_workflow(
    state: AgentState,
) -> Dict[str, Any]:
    """
    Reset workflow while keeping basic SAP configuration.
    """

    username = state.username
    tcode = state.tcode
    variant = state.variant
    plant = state.plant

    preferences = dict(
        state.preferences
    )

    state.username = username

    state.password = None

    state.tcode = tcode

    state.variant = variant

    state.plant = plant

    state.previous_period = None

    state.current_period = None

    state.previous_inventory_file = None

    state.current_inventory_file = None

    state.previous_inventory = None

    state.current_inventory = None

    state.comparison = None

    state.summary = {}

    if (
        hasattr(state, "metadata")
        and isinstance(state.metadata, dict)
    ):

        state.metadata.pop(
            "inventory_classification",
            None,
        )

        state.metadata.pop(
            "inventory_analysis",
            None,
        )

    if (
        hasattr(state, "results")
        and isinstance(state.results, dict)
    ):

        state.results.pop(
            "classification",
            None,
        )

        state.results.pop(
            "analysis",
            None,
        )

        state.results.pop(
            "inventory_analysis",
            None,
        )

    state.report_file = None

    state.report_path = None

    state.current_step = "reset"

    state.completed_steps = []

    state.errors = []

    state.warnings = []

    state.preferences = preferences

    state.status = "idle"

    return tool_result(
        True,
        "Agent workflow has been reset.",
        state=state.safe_dict(),
    )


# ============================================================
# 21. VALIDATE GENERATED REPORT
# ============================================================

def validate_report(
    state: Optional[AgentState] = None,
) -> Dict[str, Any]:
    """
    Validate the generated BSP inventory report.

    Checks:
        1. Report exists.
        2. Workbook can be opened.
        3. Required sheets exist.
        4. Classification columns exist.
    """

    import openpyxl

    project_root = (
        Path(__file__)
        .resolve()
        .parents[3]
    )

    # --------------------------------------------------------
    # DETERMINE REPORT PATH
    # --------------------------------------------------------

    report_file = None

    if state is not None:

        state_report = getattr(
            state,
            "report_file",
            None,
        )

        if state_report:

            report_file = Path(
                state_report
            )

    if report_file is None:

        report_file = (
            project_root
            / "data"
            / "reports"
            / "BSP_Inventory_Report.xlsx"
        )

    # --------------------------------------------------------
    # CHECK FILE EXISTS
    # --------------------------------------------------------

    if not report_file.exists():

        return {
            "success": False,
            "valid": False,
            "message": (
                "Generated inventory report was not found."
            ),
            "report_file": str(
                report_file
            ),
        }

    workbook = None

    try:

        # ----------------------------------------------------
        # OPEN WORKBOOK
        # ----------------------------------------------------

        workbook = openpyxl.load_workbook(
            report_file,
            read_only=True,
            data_only=False,
        )

        sheet_names = workbook.sheetnames

        # ----------------------------------------------------
        # CHECK REQUIRED SHEETS
        # ----------------------------------------------------

        expected_sheets = [
            "Inventory Comparison",
            "Summary",
        ]

        missing_sheets = [
            sheet
            for sheet in expected_sheets
            if sheet not in sheet_names
        ]

        if missing_sheets:

            return {
                "success": False,
                "valid": False,
                "message": (
                    "Generated report is missing expected sheets."
                ),
                "missing_sheets": missing_sheets,
                "available_sheets": sheet_names,
                "report_file": str(
                    report_file
                ),
            }

        # ----------------------------------------------------
        # CHECK CLASSIFICATION COLUMNS
        # ----------------------------------------------------

        comparison_sheet = workbook[
            "Inventory Comparison"
        ]

        header_row = next(
            comparison_sheet.iter_rows(
                min_row=1,
                max_row=1,
                values_only=True,
            ),
            (),
        )

        columns = {
            str(value).strip()
            for value in header_row
            if value is not None
        }

        required_columns = {
            "Inventory_Type",
            "Classification_Reason",
        }

        missing_columns = sorted(
            required_columns - columns
        )

        if missing_columns:

            return {
                "success": False,
                "valid": False,
                "message": (
                    "Generated report is missing required "
                    "RI / Capital classification columns."
                ),
                "missing_columns": missing_columns,
                "available_columns": sorted(
                    columns
                ),
                "report_file": str(
                    report_file
                ),
                "sheets": sheet_names,
            }

        # ----------------------------------------------------
        # VALIDATION SUCCESS
        # ----------------------------------------------------

        return {
            "success": True,
            "valid": True,
            "message": (
                "Inventory report validated successfully."
            ),
            "report_file": str(
                report_file
            ),
            "sheets": sheet_names,
            "classification_columns_present": True,
        }

    except Exception as exc:

        return {
            "success": False,
            "valid": False,
            "message": (
                f"Generated report could not be opened: {exc}"
            ),
            "report_file": str(
                report_file
            ),
        }

    finally:

        if workbook is not None:

            workbook.close()


# ============================================================
# 22. RI / CAPITAL CLASSIFICATION
# ============================================================

def classify_inventory_materials(
    state: AgentState,
) -> Dict[str, Any]:
    """
    Classify comparison materials using RI and Capital
    reference workbooks.

    If reference workbooks are unavailable, the classifier
    does not guess.
    """

    if state.comparison is None:

        return tool_result(
            False,
            "Inventory comparison is required before classification.",
        )

    try:

        classified_df, classification_summary = (
            classify_inventory(
                state.comparison
            )
        )

        # ----------------------------------------------------
        # STORE CLASSIFIED DATA
        # ----------------------------------------------------

        state.comparison = classified_df

        # ----------------------------------------------------
        # IMPORTANT:
        # The current classifier uses:
        #
        #     reference_files_available
        #
        # Older versions may use:
        #
        #     references_available
        #
        # Default MUST be False.
        # ----------------------------------------------------

        references_available = bool(
            classification_summary.get(
                "reference_files_available",
                classification_summary.get(
                    "references_available",
                    False,
                ),
            )
        )

        # ----------------------------------------------------
        # BUILD RESULT
        # ----------------------------------------------------

        result = {
            "success": True,

            "message": (
                "RI / Capital classification completed."
                if references_available
                else (
                    "RI / Capital classification completed "
                    "with reference files unavailable. "
                    "Materials were not guessed or forced "
                    "into RI/Capital."
                )
            ),

            "classification": (
                classification_summary
            ),

            "references_available": (
                references_available
            ),

            "comparison": (
                classified_df.to_dict(
                    orient="records"
                )
            ),
        }

        # ----------------------------------------------------
        # STORE RESULT
        # ----------------------------------------------------

        state.add_result(
            "classification",
            result,
        )

        if not hasattr(
            state,
            "metadata",
        ):

            state.metadata = {}

        state.metadata[
            "inventory_classification"
        ] = result

        # ----------------------------------------------------
        # MARK COMPLETED
        # ----------------------------------------------------

        state.mark_completed(
            "inventory_classification"
        )

        return result

    except Exception as exc:

        try:

            state.add_error(
                f"Inventory classification failed: {exc}"
            )

        except Exception:

            pass

        return {
            "success": False,
            "message": (
                f"Inventory classification failed: {exc}"
            ),
        }


# ============================================================
# 23. TOOL REGISTRY
# ============================================================

AGENT_TOOLS = {

    "get_agent_configuration":
        get_agent_configuration,

    "validate_sap_credentials":
        validate_sap_credentials,

    "login_to_sap":
        login_to_sap,

    "prepare_sap_workflow":
        prepare_sap_workflow,

    "run_mc1":
        run_mc1,

    "select_variant":
        select_variant,

    "select_report_period":
        select_report_period,

    "run_plant_analysis":
        run_plant_analysis,

    "export_inventory_from_sap":
        export_inventory_from_sap,

    "process_inventory":
        process_inventory,

    "process_previous_inventory":
        process_previous_inventory,

    "process_current_inventory":
        process_current_inventory,

    "compare_inventory_files":
        compare_inventory_files,

    "analyze_inventory":
        analyze_inventory,

    "get_comparison_summary":
        get_comparison_summary,

    "generate_report":
        generate_report,

    "get_report_path":
        get_report_path,

    "get_workflow_status":
        get_workflow_status,

    "reset_workflow":
        reset_workflow,

    "validate_report":
        validate_report,

    "classify_inventory_materials":
        classify_inventory_materials,
    
    "generate_bsp_inventory_workbook":
        generate_bsp_inventory_workbook,
}


# ============================================================
# TOOL NAMES
# ============================================================

def list_agent_tools() -> list[str]:
    """
    Return all available agent tool names.
    """

    return list(
        AGENT_TOOLS.keys()
    )